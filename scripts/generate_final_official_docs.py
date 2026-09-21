import docx, re, json, sys, os

doc_path = '/Users/vu/Desktop/Quân Vault/Quan Vault/HVNH-TÀI-LIỆU-MÔN-HỌC/THực hành nghề luật/thầy hải /Nhóm 13 - Hồ sơ pháp lý TVPAY - Điều lệ và Thỏa thuận ví điện tử - Hoàn thiện trình bày.docx'
doc = docx.Document(doc_path)
p_texts = [p.text.strip() for p in doc.paragraphs]

def table_to_md(table):
    rows = []
    for r in table.rows:
        row = [c.text.strip().replace('\n', ' ') for c in r.cells]
        rows.append(row)
    if not rows: return ""
    header = "| " + " | ".join(rows[0]) + " |"
    sep = "| " + " | ".join(["---"] * len(rows[0])) + " |"
    body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n" + "\n".join([header, sep] + body) + "\n"

md_table_nganh_nghe = table_to_md(doc.tables[2])
md_table_co_phan = table_to_md(doc.tables[3])
md_table_pl3_1 = table_to_md(doc.tables[10])
md_table_pl3_2 = table_to_md(doc.tables[11])
md_table_pl7 = table_to_md(doc.tables[12])

def clean_paras(start_p, end_p):
    res = []
    for i in range(start_p, end_p):
        t = p_texts[i]
        if not t: continue
        t = re.sub(r'^(Điều\s+\d+):', r'\1.', t)
        res.append(t)
    return res

# --- ĐIỀU LỆ CHAPTERS ---
# Chap 1
c1_paras = clean_paras(173, 208)
c1_new = []
for p in c1_paras:
    c1_new.append(p)
    if p.startswith('Điều 3.'):
        c1_new.append(md_table_nganh_nghe.strip())
c1_content = "\n\n".join(c1_new)

# Chap 2
c2_paras = clean_paras(208, 412)
c2_new = []
for p in c2_paras:
    c2_new.append(p)
    if p.startswith('Điều 6.'):
        c2_new.append(md_table_co_phan.strip())
c2_content = "\n\n".join(c2_new)

# Chap 3
c3_paras = clean_paras(412, 786)
c3_new = []
for p in c3_paras:
    if p.startswith('1.   Chủ tọa, người ghi biên bản') or p.startswith('1. Chủ tọa, người ghi biên bản'):
        p = re.sub(r'^1\.\s*', '3. ', p)
    elif p.startswith('2.   Biên bản họp Hội đồng quản trị') or p.startswith('2. Biên bản họp Hội đồng quản trị'):
        p = re.sub(r'^2\.\s*', '4. ', p)
    elif p.startswith('3.   Biên bản lập bằng tiếng Việt') or p.startswith('3. Biên bản lập bằng tiếng Việt'):
        p = re.sub(r'^3\.\s*', '5. ', p)
    c3_new.append(p)
c3_content = "\n\n".join(c3_new)

# Chap 4
c4_content = "\n\n".join(clean_paras(786, 802))

# Chap 5
c5_content = "\n\n".join(clean_paras(802, 826))

# Chap 6
c6_paras = clean_paras(826, 849)
c6_new = []
for p in c6_paras:
    if p.startswith('Công ty chia lợi nhuận cho các cổ đông'):
        p = '1. ' + p
    elif p.startswith('Trường hợp kết quả kinh doanh của năm tài chính bị lỗ'):
        p = '1. ' + p
    elif p.startswith('Sử dụng các quỹ dự trữ của Công ty để bù đắp lỗ'):
        p = 'a) ' + p
    elif p.startswith('Chuyển lỗ sang năm sau'):
        p = 'b) ' + p
    elif p.startswith('Các phương án xử lý khác'):
        p = 'c) ' + p
    elif p.startswith('Đại hội đồng cổ đông quyết định mức chi trả cổ tức'):
        p = '- ' + p
    elif p.startswith('Nếu chi trả bằng tiền mặt'):
        p = '- ' + p
    elif p.startswith('Cổ tức có thể được thanh toán bằng chuyển khoản'):
        p = '- ' + p
    elif p.startswith('Nếu chi trả cổ tức bằng cổ phần'):
        p = '- ' + p
    c6_new.append(p)
c6_content = "\n\n".join(c6_new)

# Chap 7
c7_content = "\n\n".join(clean_paras(849, 862))

# Chap 8
c8_content = "\n\n".join(clean_paras(862, 905))

# Chap 9
c9_content = "\n\n".join(clean_paras(905, 923))

dl_chapters = [
    {
        "id": "dl-chap-1",
        "number": "Chương I",
        "title": "QUY ĐỊNH CHUNG",
        "articleCount": 4,
        "summary": "Tên gọi TVPAY, trụ sở chính 89 Láng Hạ, ngành nghề 6619 và Người đại diện theo pháp luật.",
        "content": c1_content
    },
    {
        "id": "dl-chap-2",
        "number": "Chương II",
        "title": "VỐN ĐIỀU LỆ, CỔ ĐÔNG SÁNG LẬP, QUYỀN VÀ NGHĨA VỤ CỔ ĐÔNG",
        "articleCount": 25,
        "summary": "Vốn điều lệ 50 tỷ đồng của 5 cổ đông sáng lập, phân loại cổ phần, chuyển nhượng và mua lại.",
        "content": c2_content
    },
    {
        "id": "dl-chap-3",
        "number": "Chương III",
        "title": "CƠ CẤU TỔ CHỨC, QUẢN TRỊ VÀ KIỂM SOÁT",
        "articleCount": 32,
        "summary": "Thẩm quyền ĐHĐCĐ, HĐQT, Tổng Giám đốc, Ban kiểm soát và quy chế bỏ phiếu biểu quyết.",
        "content": c3_content
    },
    {
        "id": "dl-chap-4",
        "number": "Chương IV",
        "title": "CHỦ SỞ HỮU HƯỞNG LỢI CỦA CÔNG TY",
        "articleCount": 4,
        "summary": "Kê khai, xác định, cập nhật và lưu giữ thông tin chủ sở hữu hưởng lợi theo luật AML.",
        "content": c4_content
    },
    {
        "id": "dl-chap-5",
        "number": "Chương V",
        "title": "THÙ LAO, TIỀN LƯƠNG VÀ THƯỞNG CHO NGƯỜI QUẢN LÝ",
        "articleCount": 3,
        "summary": "Nguyên tắc trả thù lao HĐQT, Ban kiểm soát và tiền lương Tổng Giám đốc; công khai lợi ích liên quan.",
        "content": c5_content
    },
    {
        "id": "dl-chap-6",
        "number": "Chương VI",
        "title": "TÀI CHÍNH, PHÂN PHỐI LỢI NHUẬN VÀ LẬP QUỸ",
        "articleCount": 4,
        "summary": "Năm tài chính, tỷ lệ trích lập quỹ dự trữ 2%, quỹ phúc lợi 2%, quỹ khen thưởng 2% và xử lý lỗ.",
        "content": c6_content
    },
    {
        "id": "dl-chap-7",
        "number": "Chương VII",
        "title": "NGUYÊN TẮC GIẢI QUYẾT TRANH CHẤP NỘI BỘ",
        "articleCount": 5,
        "summary": "Trình tự thương lượng 30 ngày, hòa giải, thỏa thuận trọng tài VIAC và nghĩa vụ bảo mật.",
        "content": c7_content
    },
    {
        "id": "dl-chap-8",
        "number": "Chương VIII",
        "title": "TỔ CHỨC LẠI, GIẢI THỂ VÀ THANH LÝ TÀI SẢN",
        "articleCount": 6,
        "summary": "Điều kiện chia tách, sáp nhập, trình tự giải thể doanh nghiệp và thứ tự ưu tiên thanh toán nghĩa vụ.",
        "content": c8_content
    },
    {
        "id": "dl-chap-9",
        "number": "Chương IX",
        "title": "HIỆU LỰC THI HÀNH VÀ SỬA ĐỔI ĐIỀU LỆ",
        "articleCount": 3,
        "summary": "Hiệu lực kể từ ngày cấp phép, thể thức sửa đổi Điều lệ và chữ ký xác thực của 5 cổ đông sáng lập.",
        "content": c9_content
    }
]

# --- THỎA THUẬN SECTIONS ---
sec1_content = "\n\n".join([p_texts[i] for i in range(929, 941) if p_texts[i]])

sec2_paras = []
for i in range(941, 1016):
    t = p_texts[i]
    if not t: continue
    if t == 'Đối tượng được cung cấp dịch vụ ví điện tử:':
        t = '1. Đối tượng được cung cấp dịch vụ ví điện tử:'
    elif t.startswith('Khách hàng là cá nhân:'):
        t = 'a) ' + t
    elif t.startswith('Khách hàng là doanh nghiệp:'):
        t = 'b) ' + t
    elif t.startswith('c. Khách hàng cung cấp tài liệu'):
        t = 'c) ' + t[3:].strip()
    elif t.startswith('Nộp tiền mặt vào tài khoản đảm bảo') or t.startswith('Nhận tiền từ tài khoản') or t.startswith('Nhận tiền từ các ví'):
        t = '- ' + t
    elif t.startswith('Rút tiền ra khỏi ví') or t.startswith('Chuyển tiền đến') or t.startswith('Thanh toán cho hàng hóa'):
        t = '- ' + t
    elif t == '5. Hồ sơ mở ví điện tử':
        t = '5. Hồ sơ mở ví điện tử:'
    elif t.startswith('Trường hợp khách cá nhân là người có quốc tịch Việt Nam:'):
        t = 'a) ' + t
    elif t.startswith('Trường hợp Khách hàng cá nhân là người có Quốc tịch nước ngoài:'):
        t = 'b) ' + t
    elif t.startswith('Trường hợp Khách hàng cá nhân là người có từ hai quốc tịch trở lên:'):
        t = 'c) ' + t
    elif t.startswith('TVPAY không cung cấp dịch vụ VĐT đối với các đối tượng'):
        t = 'd) ' + t
    elif t.startswith('Việc thu thập, sử dụng, lưu trữ'):
        t = 'a) ' + t
    elif t.startswith('Thu thập Thông tin Khách hàng:'):
        t = 'b) ' + t
    elif t.startswith('Phục vụ các yêu cầu') or t.startswith('Xử lý các trường hợp') or t.startswith('Thực thi hoặc bảo vệ') or t.startswith('Duy trì mối quan hệ'):
        t = '- ' + t
    elif t.startswith('heo yêu cầu') or t.startswith('Theo yêu cầu'):
        if t.startswith('heo yêu cầu'): t = 'T' + t
        t = '- ' + t
    elif t == '9. Các hành vi bị cấm khi mở ví điện tử':
        t = '9. Các hành vi bị cấm khi mở ví điện tử:'
    elif t.startswith('Mở hoặc duy trì điện tử') or t.startswith('Mở hoặc duy trì ví điện tử') or t.startswith('Mua, bán, cho thuê') or t.startswith('Mở hộ Ví điện tử'):
        t = '- ' + t
    elif t.startswith('Tổng hạn mức giao dịch qua các tài khoản VĐT TVPAY cá nhân'):
        t = '10. ' + t
    elif t.startswith('Biểu phí dịch vụ ví điện tử được TVPAY'):
        t = '1. ' + t
    elif t.startswith('Phí Dịch vụ gồm phí theo giao dịch'):
        t = '2. ' + t
    elif t.startswith('Phí Dịch vụ sẽ được thu bằng cách'):
        t = '3. ' + t
    sec2_paras.append(t)
sec2_content = "\n\n".join(sec2_paras)

sec3_paras = []
for i in range(1016, 1131):
    t = p_texts[i]
    if not t: continue
    if t.startswith('Các giao dịch thực hiện qua Dịch vụ'):
        t = '1. ' + t
    elif t.startswith('Ứng dụng Dịch vụ Ví điện tử có sử dụng danh bạ'):
        t = '2. ' + t
    elif t.startswith('Giao dịch của Khách hàng sẽ không được Ví TVPAY xử lý'):
        t = '3. ' + t
    elif t.startswith('Không được Ngân hàng nơi Khách hàng mở thẻ'):
        t = 'a) ' + t
    elif t.startswith('Tài khoản có số dư không đủ'):
        t = 'b) ' + t
    elif t.startswith('Giao dịch vượt hạn mức'):
        t = 'c) ' + t
    elif t.startswith('Không cung cấp đầy đủ các thông tin, chứng từ'):
        t = 'd) ' + t
    elif t.startswith('Các trường hợp khác nhận thấy giao dịch có phát sinh rủi ro'):
        t = 'đ) ' + t
    elif t.startswith('Một giao dịch cụ thể bị hạn chế hoặc cấm'):
        t = 'e) ' + t
    elif t.startswith('Khách hàng không thể huỷ, thay đổi'):
        t = '1. ' + t
    elif t.startswith('Các giao dịch đã được gửi tới Ví TVPAY'):
        t = '2. ' + t
    elif t == 'Đối với khách hàng':
        t = '1. Đối với khách hàng:'
    elif t.startswith('Vi phạm quy định của Pháp luật và Vi phạm quy định về mở'):
        t = 'a) ' + t
    elif t.startswith('Thực hiện hành vi xâm phạm quyền sở hữu'):
        t = 'b) ' + t
    elif t.startswith('Phân tán các phần mềm độc hại'):
        t = 'c) ' + t
    elif t.startswith('Cung cấp thông tin sai sự thật'):
        t = 'd) ' + t
    elif t.startswith('Sử dụng tài khoản Ví để thực hiện các giao dịch cho các mục đích rửa tiền'):
        t = 'đ) ' + t
    elif t.startswith('Mua, bán, cho thuê, chuyển nhượng tài khoản Ví'):
        t = 'e) ' + t
    elif t == '2. Đối với Ví TVPAY:':
        t = '2. Đối với Ví TVPAY:'
    elif t.startswith('Cho phép khách hàng sử dụng VĐT khi chưa thu thập'):
        t = 'a) ' + t
    elif t.startswith('Cấp tín dụng cho Khách hàng'):
        t = 'b) ' + t
    elif t.startswith('Quyền của khách hàng:'):
        t = '1. Quyền của khách hàng:'
    elif t.startswith('Được sử dụng Dịch vụ Ví điện tử trong trong phạm vi'):
        t = 'a) ' + t
    elif t.startswith('Sử dụng Dịch vụ theo quy định của Ví TVPAY'):
        t = 'b) ' + t
    elif t.startswith('Sử dụng số tiền trên tài khoản VĐT TVPAY'):
        t = 'c) ' + t
    elif t.startswith('Yêu cầu Ví TVPAY thực hiện các giao dịch (nạp tiền'):
        t = 'd) ' + t
    elif t.startswith('Được yêu cầu Ví TVPAY đóng tài khoản'):
        t = 'đ) ' + t
    elif t.startswith('Chấm dứt sử dụng Dịch vụ sau khi đã thực hiện'):
        t = 'e) ' + t
    elif t.startswith('Gửi yêu cầu khiếu nại liên quan'):
        t = 'g) ' + t
    elif t.startswith('Xem, chỉnh sửa hoặc yêu cầu chỉnh sửa dữ liệu'):
        t = 'h) ' + t
    elif t.startswith('Phản đối việc xử lý dữ liệu cá nhân'):
        t = 'i) ' + t
    elif t.startswith('Yêu cầu bồi thường thiệt hại theo quy định'):
        t = 'k) ' + t
    elif t.startswith('Các quyền khác theo thỏa thuận riêng biệt'):
        t = 'l) ' + t
    elif t.startswith('Tuân thủ các quy định của Ví TVPAY và quy định của bản Thỏa thuận'):
        t = 'a) ' + t
    elif t.startswith('Cung cấp đầy đủ và chính xác các thông tin'):
        t = 'b) ' + t
    elif t.startswith('Khách hàng chịu trách nhiệm đảm bảo tính bảo mật'):
        t = 'c) ' + t
    elif t.startswith('Trong toàn bộ quá trình thiết lập, quản lý'):
        t = 'd) ' + t
    elif t.startswith('Thông báo biến động số dư tài khoản'):
        t = 'đ) ' + t
    elif t.startswith('Thông báo Thông báo mật khẩu đăng nhập') or t.startswith('Thông báo mật khẩu đăng nhập'):
        t = 'e) ' + t
    elif t == 'Quyền của ví TVPAY.' or t == 'Quyền của ví TVPAY:':
        t = '1. Quyền của ví TVPAY:'
    elif t.startswith('Cung cấp cho cơ quan Nhà nước có thẩm quyền về thông tin Khách hàng'):
        t = 'a) ' + t
    elif t.startswith('Yêu cầu Khách hàng cung cấp thông tin cá nhân'):
        t = 'b) ' + t
    elif t.startswith('Yêu cầu khách hàng thực hiện xác minh giao dịch'):
        t = 'c) ' + t
    elif t.startswith('Yêu cầu khách hàng liên kết với tài khoản'):
        t = 'd) ' + t
    elif t.startswith('Ví TVPAY chỉ chịu trách nhiệm đối với những thiệt hại'):
        t = 'đ) ' + t
    elif t.startswith('Ví TVPAY được miễn trách nhiệm đối với những thiệt hại'):
        t = 'e) ' + t
    elif t.startswith('Được quyền phong tỏa/tạm dừng/đóng/từ chối'):
        t = 'g) ' + t
    elif t.startswith('Ví TVPAY không chịu trách nhiệm với sự sai sót'):
        t = 'h) ' + t
    elif t.startswith('Từ chối yêu cầu đóng tài khoản VĐT TVPAY'):
        t = 'i) ' + t
    elif t.startswith('Quy định và áp dụng các biện pháp đảm bảo an toàn'):
        t = 'k) ' + t
    elif t.startswith('Ví TVPAY được chủ động trích (ghi Nợ)'):
        t = 'l) ' + t
    elif t.startswith('Tuân thủ đầy đủ các quy định của Pháp luật'):
        t = 'a) ' + t
    elif t.startswith('Hỗ trợ, tư vấn, xử lý yêu cầu Khách hàng'):
        t = 'b) ' + t
    elif t.startswith('Thông báo, cảnh báo để khách hàng nhận biết'):
        t = 'c) ' + t
    elif t.startswith('Tiếp nhận đề nghị tra soát, khiếu nại'):
        t = 'd) ' + t
    elif t.startswith('Thực hiện yêu cầu giao dịch của Chủ tài khoản'):
        t = 'đ) ' + t
    elif t.startswith('Ghi Có vào tài khoản VĐT TVPAY'):
        t = 'e) ' + t
    elif t.startswith('Cập nhật kịp thời các thông tin khi có thông báo'):
        t = 'g) ' + t
    elif t.startswith('Đảm bảo bảo mật các thông tin liên quan'):
        t = 'h) ' + t
    elif t.startswith('Chịu trách nhiệm về những thiệt hại do sai sót'):
        t = 'i) ' + t
    elif t.startswith('Tuân thủ Luật Phòng, chống rửa tiền'):
        t = 'k) ' + t
    elif t.startswith('Xây dựng quy trình nội bộ về mở, sử dụng'):
        t = 'l) ' + t
    elif t.startswith('Thực hiện các thủ tục nhận biết khách hàng'):
        t = 'm) ' + t
    elif t.startswith('Xử lý hoàn trả số dư còn lại cho Khách hàng'):
        t = 'n) ' + t
    sec3_paras.append(t)
sec3_content = "\n\n".join(sec3_paras)

sec4_paras = []
for i in range(1131, 1193):
    t = p_texts[i]
    if not t: continue
    if t == 'Cập nhật thông tin khách hàng':
        t = '1. Cập nhật thông tin khách hàng:'
    elif t.startswith('Trách nhiệm theo dõi hiệu lực của giấy tờ tùy thân'):
        t = 'a) ' + t
    elif t.startswith('Phương pháp cập nhật: Xóa thông tin'):
        t = 'b) ' + t
    elif t.startswith('Thay đổi, bổ sung thông tin khách hàng'):
        t = '2. Thay đổi, bổ sung thông tin khách hàng:'
    elif t.startswith('Chủ tài khoản VĐT TVPAY có trách nhiệm thông báo'):
        t = 'a) ' + t
    elif t.startswith('Thủ tục thông báo:'):
        t = 'b) ' + t
    elif t.startswith('Ví TVPAY thực hiện phong tỏa một phần hoặc toàn bộ số tiền trên tài khoản VĐT TVPAY trong các trường hợp sau:'):
        t = '1. ' + t
    elif t.startswith('Có quyết định hoặc yêu cầu bằng văn bản của cơ quan Nhà nước có thẩm quyền theo quy định'):
        t = 'a) ' + t
    elif t.startswith('Theo yêu cầu từ các đối tác của Ví TVPAY liên quan đến tranh chấp'):
        t = 'b) ' + t
    elif t.startswith('Khi Ví TVPAY phát hiện thấy có nhầm lẫn, sai sót'):
        t = 'c) ' + t
    elif t.startswith('Ví TVPAY phát hiện có dấu hiệu gian lận, vi phạm pháp luật'):
        t = 'd) ' + t
    elif t.startswith('Ví TVPAY chấm dứt phong tỏa tài khoản VĐT TVPAY khi có một trong các điều kiện sau:'):
        t = '2. ' + t
    elif t.startswith('Kết thúc thời hạn phong tỏa;'):
        t = 'a) ' + t
    elif t.startswith('Khi có quyết định hoặc yêu cầu bằng văn bản của cơ quan Nhà nước có thẩm quyền chấm dứt'):
        t = 'b) ' + t
    elif t.startswith('Ví TVPAY đã xử lý xong sai sót, nhầm lẫn'):
        t = 'c) ' + t
    elif t.startswith('Sau khi xác minh tài khoản VĐT TVPAY không có gian lận'):
        t = 'd) ' + t
    elif t.startswith('Sau khi đã giải quyết xong tranh chấp, tra soát'):
        t = 'đ) ' + t
    elif t.startswith('Khách hàng được hoàn trả tiền vào ví điện tử trong các trường hợp sau đây:'):
        t = '1. ' + t
    elif t.startswith('Có yêu cầu hoàn trả từ bên nhận khoản tiền'):
        t = 'a) ' + t
    elif t.startswith('Khách Hàng có yêu cầu tra soát, khiếu nại'):
        t = 'b) ' + t
    elif t.startswith('Giao dịch bị lỗi: Khách Hàng đã bị trừ tiền'):
        t = 'c) ' + t
    elif t.startswith('TVPAY thực hiện báo có VĐT của Khách Hàng trong vòng 05'):
        t = '2. ' + t
    elif t.startswith('Ví TVPAY thực hiện đóng tài khoản của khách hàng trong các trường hợp:'):
        t = '1. ' + t
    elif t.startswith('Có yêu cầu đóng tài khoản VĐT TVPAY của chính Chủ tài khoản'):
        t = 'a) ' + t
    elif t.startswith('Chủ tài khoản VĐT TVPAY của cá nhân bị chết, bị tuyên bố là đã chết'):
        t = 'b) ' + t
    elif t.startswith('Ví TVPAY bị thu hồi giấy phép'):
        t = 'c) ' + t
    elif t.startswith('Chủ tài khoản VĐT TVPAY vi phạm cam kết'):
        t = 'd) ' + t
    elif t.startswith('Ví TVPAY phát hiện khách hàng sử dụng giấy tờ giả'):
        t = 'đ) ' + t
    elif t.startswith('Chủ VĐT sử dụng Ví cho mục đích lừa đảo'):
        t = 'e) ' + t
    elif t.startswith('Chủ Ví chưa hoặc không liên kết VĐT'):
        t = 'g) ' + t
    elif t.startswith('Các trường hợp khác theo quy định của pháp luật và bản Thỏa thuận này.'):
        t = 'h) ' + t
    elif t.startswith('Nghiệp vụ liên quan khi đóng tài khoản ví điện tử:'):
        t = '2. ' + t
    elif t.startswith('Thực hiện xác minh thông tin chủ tài khoản'):
        t = 'a) ' + t
    elif t.startswith('Thực hiện hủy các giao dịch đã lập'):
        t = 'b) ' + t
    elif t.startswith('Ví điện tử đã đóng sẽ không được mở lại'):
        t = 'c) ' + t
    elif t.startswith('Số dư còn lại sau khi đóng tài khoản VĐT TVPAY được xử lý như sau:'):
        t = '3. ' + t
    elif t.startswith('Yêu cầu chủ tài khoản VĐT chủ động xử lý số dư'):
        t = 'a) ' + t
    elif t.startswith('Chi trả cho người được thừa kế'):
        t = 'b) ' + t
    elif t.startswith('Chi trả theo quyết định của tòa án;'):
        t = 'c) ' + t
    elif t.startswith('Ví TVPAY xử lý theo quy định của Pháp luật đối với trường hợp người thụ hưởng'):
        t = 'd) ' + t
    elif t.startswith('Trường hợp không liên lạc được chủ VĐT'):
        t = 'đ) ' + t
    elif t.startswith('Đối với ví điện tử đã đóng nhưng còn số dư'):
        t = 'e) ' + t
    elif t.startswith('Tiếp nhận tra soát, khiếu nại') and len(t) < 40:
        t = '1. Tiếp nhận tra soát, khiếu nại:'
    elif t.startswith('Tiếp nhận thông tin tra soát, khiếu nại bao gồm qua email'):
        t = 'a) ' + t
    elif t.startswith('Thời hạn khách hàng được quyền đề nghị tra soát'):
        t = 'b) ' + t
    elif t.startswith('Thời hạn xử lý đề nghị tra soát, khiếu nại không quá 30 ngày'):
        t = 'c) ' + t
    elif t.startswith('Ví TVPAY sẽ tiến hành phong tỏa VĐT khi khách hàng yêu cầu'):
        t = 'd) ' + t
    elif t.startswith('Xử lý kết quả tra soát, khiếu nại') and len(t) < 40:
        t = '2. Xử lý kết quả tra soát, khiếu nại:'
    elif t.startswith('Trong thời hạn tối đa 05 (năm) ngày làm việc'):
        t = 'a) ' + t
    elif t.startswith('Trong trường hợp hết thời hạn xử lý tra soát'):
        t = 'b) ' + t
    elif t.startswith('Trường hợp vụ việc có dấu hiệu tội phạm'):
        t = 'c) ' + t
    elif t.startswith('Trường hợp Ví TVPAY và/hoặc Tổ chức cung ứng dịch vụ thanh toán'):
        t = 'd) ' + t
    elif t.startswith('Khách hàng có thể tra cứu trực tuyến thông tin, tiến độ'):
        t = '3. ' + t
    sec4_paras.append(t)
sec4_content = "\n\n".join(sec4_paras)

sec5_paras = []
for i in range(1193, 1215):
    t = p_texts[i]
    if not t: continue
    if t.startswith('Sự Kiện Bất Khả Kháng là các sự kiện xảy ra một cách khách quan'):
        t = '1. ' + t
    elif t.startswith('Trong trường hợp Ví TVPAY xảy ra sự kiện bất khả kháng'):
        t = '2. ' + t
    elif t.startswith('Tất cả quyền sở hữu trí tuệ tồn tại trong các sản phẩm'):
        t = '1. ' + t
    elif t.startswith('Ví TVPAY có toàn quyền, bao gồm nhưng không giới hạn'):
        t = '2. ' + t
    elif t.startswith('Khách hàng đồng ý để Ví TVPAY tự do sử dụng'):
        t = '3. ' + t
    elif t.startswith('TVPAY chỉ được miễn trách nhiệm trong phạm vi thiệt hại phát sinh trực tiếp'):
        t = '1. ' + t
    elif t.startswith('TVPAY không có trách nhiệm bù đắp tổn thất, thiệt hại của Khách Hàng'):
        t = '2. ' + t
    elif t.startswith('Trừ trường hợp Hai Bên có thỏa thuận khác bằng văn bản'):
        t = '3. ' + t
    elif t.startswith('TVPAY không buộc phải thực hiện bất kỳ nghĩa vụ nào của mình'):
        t = '4. ' + t
    elif t.startswith('Nếu TVPAY (bao gồm cả các cán bộ, người lao động'):
        t = '5. ' + t
    elif t.startswith('Các thông tin về dịch vụ và Thỏa thuận mở và sử dụng VĐT'):
        t = '1. ' + t
    elif t.startswith('Khách hàng có thể gửi thông báo, yêu cầu cho TVPAY qua các kênh'):
        t = '2. ' + t
    elif t.startswith('Ví TVPAY được phép thay đổi, sửa đổi, điều chỉnh bản Thỏa thuận này'):
        t = '1. ' + t
    elif t.startswith('TVPAY thông báo rõ nội dung sửa đổi, thời điểm áp dụng'):
        t = '2. ' + t
    elif t.startswith('Khách hàng cam kết thực hiện nghiêm chỉnh các nội dung tại Thỏa thuận này'):
        t = '3. ' + t
    elif t.startswith('Trong quá trình thực hiện, Khách hàng thừa nhận các giá trị pháp lý bản gốc'):
        t = '4. ' + t
    elif t.startswith('Khi tham gia mở và sử dụng dịch vụ Ví điện tử TVPAY, Khách hàng cam kết'):
        t = '5. ' + t
    sec5_paras.append(t)
sec5_content = "\n\n".join(sec5_paras)

sec6_content = "PHẦN C\n\nCAM KẾT & XÁC NHẬN CHUNG CỦA CHỦ TÀI KHOẢN VÀ TVPAY\n\n" + "\n\n".join([p_texts[i] for i in range(1215, 1220) if p_texts[i]])

tt_chapters = [
    {
        "id": "hd-sec-1",
        "number": "PHẦN A",
        "title": "THÔNG TIN ĐỊNH DANH KHÁCH HÀNG (eKYC Cá nhân & Tổ chức)",
        "articleCount": 1,
        "summary": "Biểu mẫu thu thập họ tên, CCCD gắn chip/Hộ chiếu, sinh trắc học, GPKD, mã số thuế và người đại diện.",
        "content": sec1_content
    },
    {
        "id": "hd-sec-2",
        "number": "MỤC I",
        "title": "QUY ĐỊNH CHUNG, ĐIỀU KIỆN MỞ VÍ & BIỂU PHÍ (Điều 1 - 3)",
        "articleCount": 3,
        "summary": "Giải thích thuật ngữ, đối tượng cá nhân/doanh nghiệp, liên kết tài khoản 1:1, hạn mức 100 triệu/tháng và biểu phí.",
        "content": sec2_content
    },
    {
        "id": "hd-sec-3",
        "number": "MỤC II",
        "title": "XỬ LÝ GIAO DỊCH, GIAO DỊCH KHÔNG HỦY NGANG & HÀNH VI CẤM (Điều 4 - 8)",
        "articleCount": 5,
        "summary": "Xử lý lệnh nạp/rút/chuyển, giao dịch không thể hủy ngang một khi đã xác thực, danh mục hành vi cấm và quyền nghĩa vụ hai bên.",
        "content": sec3_content
    },
    {
        "id": "hd-sec-4",
        "number": "MỤC III",
        "title": "PHONG TỎA VÍ, HOÀN TIỀN, ĐÓNG VÍ & TRA SOÁT KHIẾU NẠI (Điều 9 - 13)",
        "articleCount": 5,
        "summary": "Căn cứ phong tỏa ví theo quyết định tư pháp, chính sách hoàn tiền trong 05 ngày, thủ tục đóng ví và thời hạn xử lý khiếu nại 30 ngày.",
        "content": sec4_content
    },
    {
        "id": "hd-sec-5",
        "number": "MỤC IV",
        "title": "BẤT KHẢ KHÁNG, MIỄN TRỪ TRÁCH NHIỆM & ĐIỀU KHOẢN CHUNG (Điều 14 - 18)",
        "articleCount": 5,
        "summary": "Sự kiện bất khả kháng, sở hữu trí tuệ phần mềm TVPAY, miễn trách nhiệm lỗi viễn thông và cơ chế sửa đổi thỏa thuận.",
        "content": sec5_content
    },
    {
        "id": "hd-sec-6",
        "number": "PHẦN C",
        "title": "CAM KẾT & XÁC NHẬN CHUNG (Hiệu lực điện tử)",
        "articleCount": 1,
        "summary": "Cam kết bảo mật thông tin, trách nhiệm pháp lý và chứng nhận điện tử giữa TVPAY và khách hàng.",
        "content": sec6_content
    }
]

# --- APPENDICES ---
dl_appendices = [
    {
        "id": "phu-luc-1",
        "number": "Phụ Lục 1",
        "title": "NỘI DUNG VÀ VAI TRÒ CỦA ĐIỀU LỆ CÔNG TY",
        "category": "role",
        "summary": "Phân tích bản chất Điều lệ theo Khoản 2 Điều 24 Luật Doanh nghiệp 2020 (sửa đổi 2025), 4 nguyên tắc soạn thảo cốt lõi và 5 vai trò sống còn đối với quản trị công ty cổ phần.",
        "content": "\n\n".join([p_texts[i] for i in range(1220, 1236) if p_texts[i]])
    },
    {
        "id": "phu-luc-2",
        "number": "Phụ Lục 2",
        "title": "GIẢI PHÁP BẢO VỆ QUYỀN VÀ LỢI ÍCH HỢP PHÁP CỦA DOANH NGHIỆP",
        "category": "protection",
        "summary": "Hệ thống 7 giải pháp bảo vệ doanh nghiệp: cân bằng quyền lợi cổ đông, phân định thẩm quyền ĐHĐCĐ - HĐQT - TGĐ, trích lập quỹ dự phòng rủi ro tài chính đặc thù fintech, kiểm soát xung đột lợi ích, bảo vệ bí mật công nghệ và quản trị khủng hoảng an ninh mạng.",
        "content": "\n\n".join([p_texts[i] for i in range(1236, 1260) if p_texts[i]])
    },
    {
        "id": "phu-luc-3",
        "number": "Phụ Lục 3",
        "title": "MA TRẬN RỦI RO PHÁP LÝ ĐỐI VỚI ĐIỀU LỆ",
        "category": "matrix",
        "summary": "Ma trận nhận diện rủi ro pháp lý Điều lệ, hệ quả tiêu cực và phương án xử lý đối chiếu với Luật Doanh nghiệp 2020 và quy định NHNN.",
        "content": "PHỤ LỤC 3. MA TRẬN RỦI RO PHÁP LÝ ĐỐI VỚI ĐIỀU LỆ CÔNG TY CỔ PHẦN TVPAY\n\n### Bảng 1: Nhận diện và Xử lý Rủi ro Cốt lõi\n" + md_table_pl3_1 + "\n\n### Bảng 2: Rà soát Hiệu lực và Pháp lý Chuyên ngành\n" + md_table_pl3_2
    }
]

tt_appendices = [
    {
        "id": "phu-luc-4",
        "number": "Phụ Lục 4",
        "title": "NỘI DUNG VÀ VAI TRÒ CỦA THỎA THUẬN VÍ ĐIỆN TỬ",
        "category": "role",
        "summary": "Cơ sở pháp lý theo Nghị định 52/2024/NĐ-CP và Thông tư 40/2024/TT-NHNN, phân tích 6 nội dung bắt buộc trong hợp đồng mẫu dịch vụ trung gian thanh toán và vai trò thiết lập quan hệ pháp lý minh bạch.",
        "content": "\n\n".join([p_texts[i] for i in range(1264, 1280) if p_texts[i]])
    },
    {
        "id": "phu-luc-5",
        "number": "Phụ Lục 5",
        "title": "GIẢI PHÁP BẢO VỆ QUYỀN VÀ LỢI ÍCH HỢP PHÁP KHI CUNG ỨNG VÍ ĐIỆN TỬ",
        "category": "protection",
        "summary": "Hệ thống 5 đối sách phòng ngừa rủi ro cho tổ chức trung gian thanh toán: chuẩn hóa eKYC theo chuẩn FIDO2, quy định giao dịch không hủy ngang, lưu giữ chứng cứ điện tử, quy trình tra soát khiếu nại minh bạch và kiểm soát tuân thủ AML liên tục.",
        "content": "\n\n".join([p_texts[i] for i in range(1280, 1290) if p_texts[i]])
    },
    {
        "id": "phu-luc-6",
        "number": "Phụ Lục 6",
        "title": "CẢNH BÁO VÀ HƯỚNG DẪN GIAO DỊCH AN TOÀN",
        "category": "safety",
        "summary": "Cẩm nang an toàn ví điện tử: 5 hành vi cấm khi mở ví, quy tắc bảo mật mã OTP/mật khẩu, cảnh báo 4 kịch bản lừa đảo công nghệ cao phổ biến (mạo danh TVPAY, giả công an, khuyến mại ảo, hack mạng xã hội) và quy trình xử lý khẩn cấp.",
        "content": "\n\n".join([p_texts[i] for i in range(1290, 1381) if p_texts[i]])
    },
    {
        "id": "phu-luc-7",
        "number": "Phụ Lục 7",
        "title": "BẢNG RÀ SOÁT VÀ HIỆU CHỈNH PHÁP LÝ",
        "category": "audit",
        "summary": "Bảng rà soát đỏ - vàng - xanh các điều khoản nhạy cảm của Thỏa thuận ví TVPAY: khắc phục bẫy hủy số dư sau 5 năm, thu hẹp giới hạn miễn trách nhiệm và chuẩn hóa chấp thuận xử lý dữ liệu cá nhân theo Nghị định 13/2023 và Luật Dữ liệu 2025.",
        "content": "PHỤ LỤC 7. BẢNG RÀ SOÁT VÀ HIỆU CHỈNH PHÁP LÝ THỎA THUẬN VÍ ĐIỆN TỬ TVPAY\n\n" + md_table_pl7
    }
]

# Write out the complete typescript file
ts_code = f"""// Dữ liệu chính thức Điều lệ TVPAY và Hợp đồng mở ví TVPAY
// Trích xuất trực tiếp từ hồ sơ gốc của nhóm dự án TVPAY (Khoa Luật - HVNH)
// Đồng bộ 100% với văn bản pháp lý chính thức, phân cấp cấu trúc chuẩn mực

export interface DocumentArticle {{
  id: string;
  articleNumber: string;
  title: string;
  fullTitle: string;
  content: string;
}}

export interface DocumentChapter {{
  id: string;
  number: string;
  title: string;
  articleCount: number;
  summary: string;
  content: string;
  articles?: DocumentArticle[];
}}

export interface DocumentAppendix {{
  id: string;
  number: string;
  title: string;
  category: 'role' | 'protection' | 'matrix' | 'safety' | 'audit';
  summary: string;
  content: string;
}}

export interface StrategicRole {{
  title: string;
  badge: string;
  highlight: string;
  description: string;
  citation: string;
  appendixRef?: string;
}}

export interface LegalProtection {{
  title: string;
  riskTitle?: string;
  protectionTier: string;
  targetRisk: string;
  riskScenario?: string;
  mechanism: string;
  solution?: string;
  contractClause: string;
  articleRef?: string;
  impactScore: number;
  appendixRef?: string;
}}

export interface OfficialDocumentData {{
  slug: string;
  officialTitle: string;
  codeName: string;
  approvedDate: string;
  legalBases: string[];
  overview: string;
  strategicRoles: StrategicRole[];
  legalProtections: LegalProtection[];
  chapters: DocumentChapter[];
  appendices: DocumentAppendix[];
}}

const articlesCache = new Map<string, DocumentArticle[]>();

export function getChapterArticles(chap: DocumentChapter): DocumentArticle[] {{
  if (articlesCache.has(chap.id)) {{
    return articlesCache.get(chap.id)!;
  }}

  const paragraphs = chap.content.split('\\n\\n');
  const articles: DocumentArticle[] = [];
  let currentArticle: {{
    id: string;
    articleNumber: string;
    title: string;
    fullTitle: string;
    parts: string[];
  }} | null = null;
  const preambleParts: string[] = [];

  for (const rawP of paragraphs) {{
    const p = rawP.trim();
    if (!p) continue;

    // Check if line starts an article/provision
    const isChapHeading = /^(CHƯƠNG\\s+[IVX]+|QUY ĐỊNH CHUNG CỦA HỢP ĐỒNG)/i.test(p);
    const isArticle =
      /^((Điều\\s+(?:\\d+|xx|[a-z0-9]+))|PHẦN\\s+[A-Z]|MỤC\\s+[IVX]+)[\\.:\\s\\-]/i.test(p) ||
      /^(Điều\\s+(?:\\d+|xx)|PHẦN\\s+[A-Z])$/i.test(p);

    if (isArticle && !isChapHeading) {{
      if (currentArticle) {{
        articles.push({{
          id: currentArticle.id,
          articleNumber: currentArticle.articleNumber,
          title: currentArticle.title,
          fullTitle: currentArticle.fullTitle,
          content: currentArticle.parts.join('\\n\\n'),
        }});
      }}

      const firstLine = p.split('\\n')[0].trim();
      const restOfP = p.slice(firstLine.length).trim();

      const match = firstLine.match(
        /^((?:Điều\\s+(?:\\d+|xx|[a-z0-9]+))|PHẦN\\s+[A-Z]|MỤC\\s+[IVX]+)[\\.:\\-]?\\s*(.*)/i,
      );
      const articleNumber = match ? match[1].trim() : firstLine;
      const title = match && match[2] ? match[2].trim() : firstLine;

      currentArticle = {{
        id: `${{chap.id}}-art-${{articles.length + 1}}`,
        articleNumber,
        title: title || articleNumber,
        fullTitle: firstLine,
        parts: restOfP ? [restOfP] : [],
      }};
    }} else {{
      if (currentArticle) {{
        currentArticle.parts.push(p);
      }} else {{
        preambleParts.push(p);
      }}
    }}
  }}

  if (currentArticle) {{
    articles.push({{
      id: currentArticle.id,
      articleNumber: currentArticle.articleNumber,
      title: currentArticle.title,
      fullTitle: currentArticle.fullTitle,
      content: currentArticle.parts.join('\\n\\n'),
    }});
  }}

  if (articles.length === 0) {{
    articles.push({{
      id: `${{chap.id}}-art-1`,
      articleNumber: chap.number,
      title: chap.title,
      fullTitle: `${{chap.number}}. ${{chap.title}}`,
      content: chap.content,
    }});
  }}

  articlesCache.set(chap.id, articles);
  return articles;
}}

export const tvpayOfficialDocs: Record<string, OfficialDocumentData> = {{
  'dieu-le-cong-ty-co-phan': {{
    slug: 'dieu-le-cong-ty-co-phan',
    officialTitle: 'ĐIỀU LỆ CÔNG TY CỔ PHẦN TVPAY',
    codeName: 'Văn kiện Quản trị Nội bộ Lõi',
    approvedDate: 'Thông qua ngày 09 tháng 09 năm 2026',
    legalBases: [
      'Luật Doanh nghiệp số 59/2020/QH14 (sửa đổi, bổ sung năm 2025)',
      'Nghị định số 52/2024/NĐ-CP về thanh toán không dùng tiền mặt',
      'Thông tư số 40/2024/TT-NHNN hướng dẫn dịch vụ trung gian thanh toán'
    ],
    overview: 'Điều lệ Công ty Cổ phần TVPAY là "Bản Hiến pháp" nội bộ cao nhất, thiết lập toàn bộ cơ cấu vốn 50 tỷ đồng của 5 cổ đông sáng lập, phân định thẩm quyền tối cao giữa ĐHĐCĐ, HĐQT và Ban Tổng Giám đốc, đảm bảo điều kiện tiên quyết để được Ngân hàng Nhà nước cấp Giấy phép cung ứng dịch vụ trung gian thanh toán.',
    strategicRoles: [
      {{
        title: 'Bản Hiến pháp Quản trị & Điều kiện Tiên quyết Cấp phép',
        badge: 'Cấp phép NHNN',
        highlight: 'Vốn thực góp 50.000.000.000 VNĐ',
        description: 'Khoản 2 Điều 22 Nghị định 52/2024/NĐ-CP bắt buộc tổ chức cung ứng dịch vụ ví điện tử phải có vốn điều lệ thực góp tối thiểu 50 tỷ đồng. Điều lệ TVPAY xác lập chuẩn xác mệnh giá cổ phần, cam kết góp vốn bằng tiền mặt của 5 cổ đông sáng lập và cơ cấu sở hữu minh bạch.',
        citation: 'Điều 5 & Điều 6 Điều lệ TVPAY',
        appendixRef: 'phu-luc-1'
      }},
      {{
        title: 'Phân định Ranh giới Quyền lực & Chống Lạm quyền',
        badge: 'Trật tự Quản trị',
        highlight: 'Tách bạch Sở hữu và Điều hành',
        description: 'Thiết lập cơ chế kiểm soát chéo giữa Đại hội đồng cổ đông (Đại diện vốn), Hội đồng quản trị (Định hướng chiến lược) và Tổng Giám đốc (Người đại diện theo pháp luật điều hành). Mọi giao dịch vượt hạn mức hoặc có xung đột lợi ích bắt buộc phải được chấp thuận trước.',
        citation: 'Chương III & Điều 33 Điều lệ TVPAY',
        appendixRef: 'phu-luc-1'
      }},
      {{
        title: 'Bảo đảm Tính Ổn định & Tồn tại Liên tục',
        badge: 'Cam kết 40 năm',
        highlight: 'Thời hạn hoạt động 40 năm',
        description: 'Quy định rõ thời hạn hoạt động 40 năm kể từ ngày cấp phép, tạo cơ sở pháp lý vững chắc cho các đối tác ngân hàng thương mại yên tâm mở tài khoản bảo đảm thanh toán và kết nối cổng API thanh toán lâu dài.',
        citation: 'Điều 2 Khoản 3 Điều lệ TVPAY',
        appendixRef: 'phu-luc-1'
      }},
      {{
        title: 'Cơ chế Bảo vệ Bí quyết Công nghệ & Tài sản Số',
        badge: 'Tài sản Trí tuệ',
        highlight: 'Mã nguồn & Hệ thống ví thuộc về TVPAY',
        description: 'Khẳng định toàn bộ bản quyền phần mềm, mã nguồn ứng dụng Ví TVPAY, cơ sở dữ liệu khách hàng và thuật toán định danh thuộc quyền sở hữu tuyệt đối của Công ty, ngăn chặn nguy cơ bị cổ đông hoặc nhân sự công nghệ tách lập nền tảng riêng.',
        citation: 'Chương II & Phụ lục 2 Điều lệ TVPAY',
        appendixRef: 'phu-luc-2'
      }}
    ],
    legalProtections: [
      {{
        title: 'Khóa Chặt Tỷ Lệ Sở Hữu và Chuyển Nhượng Cổ Phần Sáng Lập',
        riskTitle: 'Khóa Chặt Tỷ Lệ Sở Hữu và Chuyển Nhượng Cổ Phần Sáng Lập',
        protectionTier: 'Lá chắn cốt lõi',
        targetRisk: 'Bị thâu tóm thù địch hoặc cổ đông sáng lập thoái vốn sớm làm mất tính liên tục hoạt động',
        riskScenario: 'Bị thâu tóm thù địch hoặc cổ đông sáng lập thoái vốn sớm làm mất tính liên tục hoạt động',
        mechanism: 'Ràng buộc cổ đông sáng lập chỉ được chuyển nhượng cổ phần phổ thông cho người khác không phải cổ đông sáng lập nếu được sự chấp thuận của Đại hội đồng cổ đông trong 03 năm đầu.',
        solution: 'Ràng buộc cổ đông sáng lập chỉ được chuyển nhượng cổ phần phổ thông cho người khác không phải cổ đông sáng lập nếu được sự chấp thuận của Đại hội đồng cổ đông trong 03 năm đầu.',
        contractClause: 'Điều 18 & Điều 21 Điều lệ TVPAY',
        articleRef: 'Điều 18 & Điều 21 Điều lệ TVPAY',
        impactScore: 98,
        appendixRef: 'phu-luc-2'
      }},
      {{
        title: 'Cơ Chế Trích Lập Quỹ Dự Trữ Rủi Ro Thanh Toán Bắt Buộc',
        riskTitle: 'Cơ Chế Trích Lập Quỹ Dự Trữ Rủi Ro Thanh Toán Bắt Buộc',
        protectionTier: 'Kiểm soát rủi ro',
        targetRisk: 'Mất cân đối dòng tiền khi đối mặt sự cố gian lận công nghệ hoặc yêu cầu bồi hoàn tức thời',
        riskScenario: 'Mất cân đối dòng tiền khi đối mặt sự cố gian lận công nghệ hoặc yêu cầu bồi hoàn tức thời',
        mechanism: 'Bắt buộc trích lập Quỹ dự trữ bắt buộc 2% lợi nhuận sau thuế hàng năm để đảm bảo khả năng thanh toán độc lập ngoài tài khoản đảm bảo tại ngân hàng hợp tác.',
        solution: 'Bắt buộc trích lập Quỹ dự trữ bắt buộc 2% lợi nhuận sau thuế hàng năm để đảm bảo khả năng thanh toán độc lập ngoài tài khoản đảm bảo tại ngân hàng hợp tác.',
        contractClause: 'Điều 70 & Điều 71 Điều lệ TVPAY',
        articleRef: 'Điều 70 & Điều 71 Điều lệ TVPAY',
        impactScore: 94,
        appendixRef: 'phu-luc-2'
      }},
      {{
        title: 'Kiểm Soát Xung Đột Lợi Ích & Giao Dịch Với Người Có Liên Quan',
        riskTitle: 'Kiểm Soát Xung Đột Lợi Ích & Giao Dịch Với Người Có Liên Quan',
        protectionTier: 'Lá chắn cốt lõi',
        targetRisk: 'Ban điều hành rút ruột công ty qua hợp đồng dịch vụ công nghệ hoặc thầu ngoài giá cao',
        riskScenario: 'Ban điều hành rút ruột công ty qua hợp đồng dịch vụ công nghệ hoặc thầu ngoài giá cao',
        mechanism: 'Mọi hợp đồng có giá trị từ 35% tổng tài sản hoặc giao dịch với người có liên quan phải được ĐHĐCĐ hoặc HĐQT chấp thuận bằng văn bản; người có liên quan không được quyền biểu quyết.',
        solution: 'Mọi hợp đồng có giá trị từ 35% tổng tài sản hoặc giao dịch với người có liên quan phải được ĐHĐCĐ hoặc HĐQT chấp thuận bằng văn bản; người có liên quan không được quyền biểu quyết.',
        contractClause: 'Điều 33 & Điều 67 Điều lệ TVPAY',
        articleRef: 'Điều 33 & Điều 67 Điều lệ TVPAY',
        impactScore: 96,
        appendixRef: 'phu-luc-2'
      }},
      {{
        title: 'Thỏa Thuận Trọng Tài Thương Mại VIAC & Nghĩa Vụ Bảo Mật',
        protectionTier: 'Miễn trừ trách nhiệm',
        targetRisk: 'Tranh chấp nội bộ bị đưa ra tòa công khai làm lộ bí mật kinh doanh và sụt giảm uy tín ví',
        mechanism: 'Quy định mọi tranh chấp nội bộ phải trải qua hòa giải kín 30 ngày. Nếu không thành, đưa ra Trung tâm Trọng tài Quốc tế Việt Nam (VIAC) tại Hà Nội, xét xử bằng phiên kín bảo mật tuyệt đối.',
        contractClause: 'Điều 73, Điều 74 & Điều 76 Điều lệ TVPAY',
        impactScore: 92,
        appendixRef: 'phu-luc-2'
      }}
    ],
    chapters: {json.dumps(dl_chapters, ensure_ascii=False, indent=6)},
    appendices: {json.dumps(dl_appendices, ensure_ascii=False, indent=6)}
  }},

  'mo-va-su-dung-vi-dien-tu': {{
    slug: 'mo-va-su-dung-vi-dien-tu',
    officialTitle: 'THỎA THUẬN MỞ, QUẢN LÝ VÀ SỬ DỤNG VÍ ĐIỆN TỬ TVPAY',
    codeName: 'Hợp Đồng Cung Ứng Dịch Vụ Người Dùng Lõi',
    approvedDate: 'Áp dụng chính thức từ ngày cấp phép hoạt động',
    legalBases: [
      'Nghị định 52/2024/NĐ-CP về thanh toán không dùng tiền mặt',
      'Thông tư 40/2024/TT-NHNN hướng dẫn trung gian thanh toán',
      'Luật Giao dịch điện tử số 20/2023/QH15',
      'Luật Bảo vệ quyền lợi người tiêu dùng số 19/2023/QH15',
      'Nghị định 13/2023/NĐ-CP về bảo vệ dữ liệu cá nhân'
    ],
    overview: 'Thỏa thuận mở, quản lý và sử dụng ví điện tử TVPAY là hợp đồng điện tử theo mẫu chuẩn mực quốc tế, xác lập cơ chế liên kết tài khoản ngân hàng bảo đảm 1:1, hạn mức thanh toán 100 triệu VNĐ/tháng, nguyên tắc giao dịch không hủy ngang một khi đã xác thực OTP/sinh trắc học, và các điều khoản miễn trách nhiệm sự cố kỹ thuật đối tác.',
    strategicRoles: [
      {{
        title: 'Thiết Lập Cơ Chế Pháp Lý Liên Kết Tài Khoản 1:1 Bắt Buộc',
        badge: 'Tuân thủ NĐ 52',
        highlight: 'Liên kết định danh ngân hàng hợp tác',
        description: 'Đáp ứng Điều 25 Nghị định 52/2024/NĐ-CP: khách hàng bắt buộc phải liên kết ví với tài khoản thanh toán hoặc thẻ ghi nợ chính chủ trước khi kích hoạt tính năng thanh toán, ngăn chặn triệt để dòng tiền ảo không được kiểm soát.',
        citation: 'Điều 2 Khoản 2 Thỏa thuận TVPAY',
        appendixRef: 'phu-luc-4'
      }},
      {{
        title: 'Xác Lập Nguyên Tắc Giao Dịch Không Hủy Ngang Sau Xác Thực',
        badge: 'Bảo Vệ Tính Toàn Vẹn',
        highlight: 'OTP & Sinh trắc học là chữ ký điện tử',
        description: 'Mọi lệnh chuyển tiền, thanh toán đã xác thực thành công bằng OTP hoặc dữ liệu sinh trắc học có giá trị pháp lý ràng buộc tuyệt đối. Khách hàng không thể tự ý hủy giao dịch hoặc thoái thác trách nhiệm dân sự.',
        citation: 'Điều 5 Thỏa thuận TVPAY',
        appendixRef: 'phu-luc-4'
      }},
      {{
        title: 'Hạn Mức Giao Dịch Tối Đa 100 Triệu Đồng/Tháng/Khách Hàng',
        badge: 'Phòng Ngừa AML',
        highlight: 'Kiểm soát hạn mức theo Thông tư 40',
        description: 'Áp dụng nghiêm ngặt hạn mức tổng giao dịch thanh toán qua các ví cá nhân của một khách hàng không vượt quá 100.000.000 VNĐ/tháng, bảo vệ hệ thống khỏi các hành vi rửa tiền và giao dịch bất hợp pháp quy mô lớn.',
        citation: 'Điều 2 Khoản 10 Thỏa thuận TVPAY',
        appendixRef: 'phu-luc-4'
      }},
      {{
        title: 'Cơ Chế Phong Tỏa & Hoàn Tiền Xử Lý Trong 05 Ngày Làm Việc',
        badge: 'Bảo Đảm Minh Bạch',
        highlight: 'Quy trình tra soát chuẩn mực',
        description: 'Quy định rõ ràng căn cứ phong tỏa ví theo yêu cầu tư pháp hoặc khi phát hiện lỗi ghi có nhầm lẫn; cam kết hoàn tiền vào ví trong vòng 05 ngày làm việc khi xác định lỗi kỹ thuật hệ thống, tạo niềm tin cho người dùng.',
        citation: 'Điều 10, 11 & 13 Thỏa thuận TVPAY',
        appendixRef: 'phu-luc-5'
      }}
    ],
    legalProtections: [
      {{
        title: 'Miễn Trừ Trách Nhiệm Sự Cố Hạ Tầng Ngân Hàng & Viễn Thông',
        riskTitle: 'Miễn Trừ Trách Nhiệm Sự Cố Hạ Tầng Ngân Hàng & Viễn Thông',
        protectionTier: 'Miễn trừ trách nhiệm',
        targetRisk: 'Khách hàng kiện TVPAY đòi bồi thường thiệt hại khi cổng ngân hàng đối tác hoặc mạng viễn thông nghẽn/lỗi',
        riskScenario: 'Khách hàng kiện TVPAY đòi bồi thường thiệt hại khi cổng ngân hàng đối tác hoặc mạng viễn thông nghẽn/lỗi',
        mechanism: 'TVPAY được miễn hoàn toàn trách nhiệm bồi thường đối với các thiệt hại gián tiếp phát sinh từ sự cố kỹ thuật của bên thứ ba độc lập hoặc các sự kiện bất khả kháng ngoài tầm kiểm soát hợp lý.',
        solution: 'TVPAY được miễn hoàn toàn trách nhiệm bồi thường đối với các thiệt hại gián tiếp phát sinh từ sự cố kỹ thuật của bên thứ ba độc lập hoặc các sự kiện bất khả kháng ngoài tầm kiểm soát hợp lý.',
        contractClause: 'Điều 14 & Điều 16 Khoản 1-2 Thỏa thuận TVPAY',
        articleRef: 'Điều 14 & Điều 16 Khoản 1-2 Thỏa thuận TVPAY',
        impactScore: 97,
        appendixRef: 'phu-luc-5'
      }},
      {{
        title: 'Quyền Đơn Phương Phong Tỏa Ví Khi Phát Hiện Dấu Hiệu Gian Lận',
        riskTitle: 'Quyền Đơn Phương Phong Tỏa Ví Khi Phát Hiện Dấu Hiệu Gian Lận',
        protectionTier: 'Lá chắn cốt lõi',
        targetRisk: 'Kẻ gian sử dụng ví TVPAY làm kênh trung chuyển tiền lừa đảo công nghệ cao',
        riskScenario: 'Kẻ gian sử dụng ví TVPAY làm kênh trung chuyển tiền lừa đảo công nghệ cao',
        mechanism: 'TVPAY có quyền tạm ngưng giao dịch hoặc phong tỏa số dư ví ngay lập tức khi hệ thống giám sát rủi ro phát hiện giao dịch bất thường mà không cần có quyết định trước của cơ quan công an.',
        solution: 'TVPAY có quyền tạm ngưng giao dịch hoặc phong tỏa số dư ví ngay lập tức khi hệ thống giám sát rủi ro phát hiện giao dịch bất thường mà không cần có quyết định trước của cơ quan công an.',
        contractClause: 'Điều 10 Khoản 1 Điểm d Thỏa thuận TVPAY',
        articleRef: 'Điều 10 Khoản 1 Điểm d Thỏa thuận TVPAY',
        impactScore: 99,
        appendixRef: 'phu-luc-5'
      }},
      {{
        title: 'Chấm Dứt Và Đóng Ví Khi Khách Hàng Vi Phạm Thỏa Thuận Mở Ví',
        protectionTier: 'Kiểm soát rủi ro',
        targetRisk: 'Khách hàng nặc danh, dùng giấy tờ giả mạo hoặc thuê mượn tài khoản vi phạm pháp luật',
        mechanism: 'TVPAY có quyền đơn phương chấm dứt hợp đồng, đóng ví vĩnh viễn và chuyển hồ sơ sang cơ quan công an khi phát hiện khách hàng cung cấp thông tin eKYC sai lệch hoặc mở hộ tài khoản cho người khác.',
        contractClause: 'Điều 6, Điều 12 Khoản 1 Thỏa thuận TVPAY',
        impactScore: 95,
        appendixRef: 'phu-luc-5'
      }},
      {{
        title: 'Bảo Vệ Quyền Xử Lý Dữ Liệu Cá Nhân & Miễn Trừ Trách Nhiệm OTP',
        protectionTier: 'Lá chắn cốt lõi',
        targetRisk: 'Khách hàng bị lừa lộ OTP cho kẻ mạo danh rồi quay lại khiếu nại TVPAY để lộ lọt thông tin',
        mechanism: 'Quy định khách hàng chịu trách nhiệm bảo mật tuyệt đối thiết bị, mật khẩu và mã OTP. Mọi giao dịch xác thực đúng mã OTP được coi là ý chí của chính khách hàng và TVPAY không có nghĩa vụ hoàn tiền.',
        contractClause: 'Điều 7 Khoản 2 & Điều 8 Khoản 1 Điểm h Thỏa thuận TVPAY',
        impactScore: 96,
        appendixRef: 'phu-luc-5'
      }}
    ],
    chapters: {json.dumps(tt_chapters, ensure_ascii=False, indent=6)},
    appendices: {json.dumps(tt_appendices, ensure_ascii=False, indent=6)}
  }}
}};
"""

target_file = 'src/lib/tvpay-official-docs.ts'
with open(target_file, 'w', encoding='utf-8') as f:
    f.write(ts_code)

print(f"File {target_file} written successfully ({len(ts_code)} bytes).")
